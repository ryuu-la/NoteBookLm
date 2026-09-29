import pytest

from local_notebook.ingestion.parsers import parse
from local_notebook.ingestion.jobs import add_file


def test_csv_retains_headers_and_row_locations(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("name,score\nAda,95\nGrace,98", encoding="utf-8")
    blocks = list(parse(path))
    assert "name | score" in blocks[0].text
    assert "rows 2–3" in blocks[0].locator


def test_docx_paragraphs_and_tables(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("A useful paragraph.")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Table value"
    path = tmp_path / "document.docx"
    document.save(path)
    blocks = list(parse(path))
    assert {block.locator for block in blocks} == {"Paragraph 1", "Table 1"}


def test_pptx_slide_locations(tmp_path):
    from pptx import Presentation
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Retrieval pipeline"
    path = tmp_path / "slides.pptx"
    slides.save(path)
    blocks = list(parse(path))
    assert blocks[0].locator == "Slide 1"
    assert "Retrieval pipeline" in blocks[0].text


def test_xlsx_sheet_and_rows(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Grades"
    sheet.append(["Student", "Score"])
    sheet.append(["Ada", 99])
    path = tmp_path / "grades.xlsx"
    book.save(path)
    block = list(parse(path))[0]
    assert "Grades" in block.locator and "99" in block.text


def test_pdf_page_citations(tmp_path):
    from fpdf import FPDF
    document = FPDF()
    for index in range(3):
        document.add_page()
        document.set_font("Helvetica", size=12)
        document.cell(0, 10, f"Important evidence on the document page number {index + 1}.")
    path = tmp_path / "document.pdf"
    document.output(path)
    assert [block.locator for block in parse(path)] == ["Page 1", "Page 2", "Page 3"]


def test_html_removes_scripts(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<script>steal()</script><h1>Useful evidence</h1>")
    assert "steal" not in list(parse(path))[0].text


def test_unknown_format_fails_explicitly(library):
    with pytest.raises(ValueError, match="Unsupported"):
        add_file(library, "app.exe", b"MZ")


def test_filename_cannot_escape_storage(library):
    from local_notebook import storage as db
    source = add_file(library, "../../secret.txt", b"A harmless test")
    assert db.one("SELECT name FROM sources WHERE id=?", (source,))["name"] == "secret.txt"


def test_pdf_fallback_retains_text_and_progress(tmp_path, monkeypatch):
    import sys
    from fpdf import FPDF
    document = FPDF()
    document.set_font("Helvetica", size=12)
    for number in range(2):
        document.add_page()
        document.cell(0, 10, f"Readable evidence on page {number + 1}.")
    path = tmp_path / "fallback.pdf"
    document.output(path)
    monkeypatch.setitem(sys.modules, "pypdfium2", None)
    blocks = list(parse(path))
    assert [block.progress for block in blocks] == [.5, 1.0]
    assert "Readable evidence on page 2" in blocks[-1].text


def test_encrypted_pdf_has_an_actionable_error(tmp_path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.encrypt("test-password")
    path = tmp_path / "locked.pdf"
    writer.write(path)
    with pytest.raises(ValueError, match="password protected"):
        list(parse(path))
