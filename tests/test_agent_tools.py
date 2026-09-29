import asyncio
import io
import json

import pytest
from openpyxl import load_workbook

from local_notebook import agent, data_tools, research, storage as db
from local_notebook.ingestion import jobs
from local_notebook.retrieval.search import Evidence


def table():
    return {'title': 'Sales', 'columns': ['Region', 'Sales'],
            'rows': [['East', 10], ['West', 20], ['East', 30], ['East', None]],
            'coverage': 'All rows', 'citations': []}


def test_full_source_calculations_and_exports(library):
    source = jobs.add_file(library, 'sales.csv', b'Region,Sales\nEast,10\nWest,20\nEast,30\nEast,\n')
    db.execute("UPDATE sources SET status='ready' WHERE id=?", (source,))
    tables = data_tools.source_tables(library)
    result = data_tools.apply_recipe(tables[0], data_tools.Recipe(
        title='Sales by region', table=0, group_by='Region', value='Sales', operation='sum', chart='bar'))
    assert result['rows'] == [['East', 40], ['West', 20]]
    assert result['chart']['values'] == [40, 20]
    book = load_workbook(io.BytesIO(data_tools.spreadsheet_bytes(result)), data_only=False)
    assert list(book['Data'].values)[1:] == [('East', 40), ('West', 20)]
    assert book['Data'].freeze_panes == 'A2'
    assert book['Sources']['B2'].value == 'sales.csv'
    assert 'East,40' in data_tools.spreadsheet_bytes(result, 'csv').decode('utf-8-sig')


def test_empty_cells_and_unit_mismatch():
    data = table()
    result = data_tools.apply_recipe(data, data_tools.Recipe(
        title='Average', table=0, group_by='Region', value='Sales', operation='mean'))
    assert result['rows'][0] == ['East', 20]
    data['rows'] = [['East', '$1 million'], ['West', '$500 thousand']]
    result = data_tools.apply_recipe(data, data_tools.Recipe(title='Revenue', table=0, chart='bar', x='Region', y='Sales'))
    assert result['chart']['values'] == [1000000, 500000]
    assert result['chart']['unit'] == '$'
    data['rows'][1][1] = '€500 thousand'
    with pytest.raises(ValueError, match='mixed'):
        data_tools.apply_recipe(data, data_tools.Recipe(title='Bad sum', table=0, operation='sum', value='Sales'))


def test_spreadsheet_formula_injection():
    data = table()
    data['rows'] = [['=HYPERLINK("https://example.org")', 12], ['@SUM(1)', -4]]
    book = load_workbook(io.BytesIO(data_tools.spreadsheet_bytes(data)), data_only=False)
    assert book['Data']['A2'].data_type == 's'
    assert book['Data']['A2'].value.startswith("'=")
    assert book['Data']['B3'].value == -4


def test_no_silent_partial_totals(monkeypatch):
    monkeypatch.setattr(data_tools, 'MAX_ROWS', 1)
    with pytest.raises(ValueError, match='no partial totals'):
        data_tools.make_table('Large', [['x'], [1], [2]], {})


def test_extraction_rejects_fabricated_values(monkeypatch):
    async def complete(*args, **kwargs):
        return json.dumps({'title': 'Sales', 'columns': ['Sales'],
                           'rows': [{'values': ['12'], 'citation': 1, 'quote': 'Sales were 1200.'}]})
    monkeypatch.setattr(data_tools.providers, 'complete', complete)
    evidence = Evidence([{'name': 'Sales', 'locator': 'Page 1', 'text': 'Sales were 1200.'}], 'test', 0)
    with pytest.raises(ValueError, match='not present'):
        asyncio.run(data_tools.extract_table('Sales', evidence))


def test_avatar_parallel_tools_share_data_and_use_quick_search(library, monkeypatch):
    active, peak, prepared, depths = 0, 0, 0, []

    async def complete(*args, **kwargs):
        return json.dumps({'tools': ['analytics', 'spreadsheet', 'deep_research'],
                           'queries': ['Avatar cast production', 'Avatar box office']})

    async def search(book, question, history, on_status, *, depth):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        depths.append(depth)
        await asyncio.sleep(.01)
        active -= 1
        return Evidence([{'id': question, 'name': question, 'locator': 'https://example.org',
                          'kind': 'web', 'url': 'https://example.org', 'text': 'Avatar earned 100.'}], 'test', 1)

    async def prepare(*args):
        nonlocal prepared
        prepared += 1
        return {**table(), 'statistics': {}, 'chart': None}

    monkeypatch.setattr(agent.providers, 'complete', complete)
    monkeypatch.setattr(agent.research, 'research', search)
    monkeypatch.setattr(agent, 'source_tables', lambda book: [])
    monkeypatch.setattr(agent, 'prepare_data', prepare)
    result = asyncio.run(agent.run(library, 'Quick search Avatar cast and production; chart sales and create spreadsheet', web=True))
    assert peak == 2 and depths == ['quick', 'quick']
    assert prepared == 1 and len(result.outputs) == 2
    assert not result.deep
    user = db.add_message(library, 'user', 'Question')
    identifier = agent.commit_answer(library, result, 'Answer [1]', [])
    blocks = json.loads(db.one('SELECT content FROM message_blocks WHERE message_id=?', (identifier,))['content'])
    assert [b['kind'] for b in blocks] == ['analytics', 'spreadsheet', 'text']
    assert len(db.rows('SELECT * FROM artifacts WHERE notebook_id=?', (library,))) == 2
    db.replace_turn(library, user, 'Retry', 'New answer')
    assert not db.rows('SELECT * FROM message_blocks')
    assert len(db.rows('SELECT * FROM artifacts')) == 2


def test_cancel_propagates_to_parallel_tasks(library, monkeypatch):
    cancelled = []
    async def complete(*args, **kwargs):
        return '{"tools":[],"queries":["one","two"]}'
    async def search(*args, **kwargs):
        try:
            await asyncio.sleep(10)
        except asyncio.CancelledError:
            cancelled.append(True)
            raise
    monkeypatch.setattr(agent.providers, 'complete', complete)
    monkeypatch.setattr(agent.research, 'research', search)
    async def scenario():
        task = asyncio.create_task(agent.run(library, 'Search', web=True))
        await asyncio.sleep(.02)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    asyncio.run(scenario())
    assert len(cancelled) == 2
    assert not db.rows('SELECT * FROM artifacts')


def test_tool_commit_rolls_back_revision_on_failure(library, monkeypatch):
    user = db.add_message(library, 'user', 'Original')
    db.add_message(library, 'assistant', 'Keep this')
    def fail(*args, **kwargs):
        raise ValueError('Storage failed')
    monkeypatch.setattr(agent, 'save_outputs', fail)
    with pytest.raises(ValueError):
        agent.commit_answer(library, agent.Run(Evidence([], 'test', 0)), 'New', [], user_id=user, question='Edited')
    assert [m['text'] for m in db.messages(library)] == ['Original', 'Keep this']


def test_routing():
    assert research.wants_web('Quick search about Avatar movie cast')
    assert 'deep_research' not in agent.requested_tools('Quick search, not deep research, and create charts')
    assert agent.requested_tools('Deep research the market and create a spreadsheet') == ['spreadsheet', 'deep_research']
    assert not research.use_web('Do not search online; use my data', 'web')


def test_inline_mindmap_and_weekly_chart(library, monkeypatch):
    question = 'Create a proper mind map and proper grpah for learning expand it into weeks'
    assert set(agent.requested_tools(question)) == {'mindmap', 'analytics'}
    async def complete(*args, **kwargs):
        return '{"tools":["mindmap","analytics"],"queries":[]}'
    async def search(*args, **kwargs):
        return Evidence([{'id': 'p1', 'name': 'Course', 'locator': 'Page 1', 'text': 'Learn APIs and testing.'}], 'test', 0)
    async def generate(*args, **kwargs):
        assert kwargs['persist'] is False
        return {'title': 'Weekly plan', 'root': {'name': 'Course', 'children': [
            {'name': 'Week 1: APIs', 'children': [{'name': 'Routing'}, {'name': 'Validation'}]},
            {'name': 'Week 2: Testing', 'children': [{'name': 'Tests'}]}]}, 'citations': []}
    monkeypatch.setattr(agent.providers, 'complete', complete)
    monkeypatch.setattr(agent.research, 'research', search)
    monkeypatch.setattr(agent, 'generate', generate)
    monkeypatch.setattr(agent, 'source_tables', lambda book: [])
    result = asyncio.run(agent.run(library, question, web=True))
    assert [o['kind'] for o in result.outputs] == ['mindmap', 'analytics']
    assert result.outputs[1]['data']['chart']['values'] == [2, 1]
    answer = 'Introduction\n\n[[tool:mindmap]]\n\nWeekly workload\n\n[[tool:analytics]]\n\nConclusion'
    message = agent.commit_answer(library, result, answer, [])
    blocks = json.loads(db.one('SELECT content FROM message_blocks WHERE message_id=?', (message,))['content'])
    rendered = agent.interleaved_answer(answer, blocks)
    assert [b['kind'] for b in rendered[:5]] == ['answer_text', 'mindmap', 'answer_text', 'analytics', 'answer_text']
    assert 'Conclusion' in rendered[4]['text']
    assert agent.interleaved_answer('Intro\n\nEnd', blocks)[1]['kind'] == 'mindmap'
    assert len(db.rows("SELECT id FROM artifacts WHERE kind IN ('mindmap','analytics')")) == 2


def test_lifespan_chart_preserves_source_ranges():
    data = {'columns': ['Philosopher', 'Lifespan'],
            'rows': [['Hegel', '1770–1831'], ['Descartes', '1596–1650'], ['Hobbes', '1588–1679']]}
    chart = data_tools.suggested_chart(data)
    assert chart['horizontal'] and chart['values'] == [61, 54, 91]
    assert chart['unit'] == 'Approximate years lived'
    assert data['rows'][0][1] == '1770–1831'
    data['rows'][0][1] = 'Unknown'
    assert data_tools.suggested_chart(data) is None


def test_default_numeric_chart_and_non_numeric_fallback():
    assert data_tools.suggested_chart({'columns': ['Region', 'Sales'], 'rows': [['East', 10], ['West', 20]]})['values'] == [10, 20]
    assert data_tools.suggested_chart({'columns': ['Name', 'Belief'], 'rows': [['Hegel', 'Idealism']]}) is None


def test_deep_research_partial_failure_and_report(library, monkeypatch):
    async def complete(*args, **kwargs):
        return '{"tools":["deep_research"],"queries":["working","unavailable"]}'
    async def search(book, query, history, status, *, depth):
        assert depth == 'deep'
        if query == 'unavailable':
            raise ValueError('Website unavailable')
        return Evidence([{'id': 'web-1', 'kind': 'web', 'name': 'Report evidence', 'locator': 'https://example.org',
                          'url': 'https://example.org', 'text': 'Supported finding.'}], 'Agent', 1)
    monkeypatch.setattr(agent.providers, 'complete', complete)
    monkeypatch.setattr(agent.research, 'research', search)
    result = asyncio.run(agent.run(library, 'Deep research this topic', web=True))
    assert result.deep and result.errors and len(result.evidence.passages) == 1
    assert '1500–2500' in result.answer_instruction()
    ids = agent.save_outputs(library, result, '# Report\nSupported finding. [1]')
    saved = db.one('SELECT * FROM artifacts WHERE id=?', (ids[0],))
    assert saved['kind'] == 'deep_research'
    assert json.loads(saved['citations'])[0]['url'] == 'https://example.org'
