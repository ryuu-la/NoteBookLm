"""Index migration, visual provenance and retrieval isolation (model-free fixtures)."""
import asyncio
import json
from io import BytesIO

import httpx
import numpy as np
from PIL import Image

from local_notebook import chat, providers, storage as db
from local_notebook.config import EMBED_DIM, INDEX_VERSION
from local_notebook.ingestion import jobs
from local_notebook.ingestion.media import answer_images, open_image
from local_notebook.ingestion.parsers import parse
from local_notebook.retrieval import vectors
from local_notebook.retrieval.embedding_cache import embed_cached
from local_notebook.retrieval.search import Evidence, retrieve


def test_search_overtakes_indexing_and_table_access_does_not_wait(library):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    import time

    entered, release = Event(), Event()
    calls = []
    class Model:
        def encode(self, inputs, **kwargs):
            calls.append(inputs[0])
            if inputs[0] == 'document-one':
                entered.set()
                assert release.wait(5)
            output = np.zeros((len(inputs), EMBED_DIM), dtype=np.float32)
            output[:, 0] = 1
            return output
    model = object.__new__(vectors.GemmaEmbedding)
    model.model = Model()
    with ThreadPoolExecutor(max_workers=3) as pool:
        indexing = pool.submit(model.encode, ['document-one', 'document-two'], prompt_name='Document', batch_size=1)
        try:
            assert entered.wait(5)
            query = pool.submit(model.query_embed, 'question')
            deadline = time.monotonic() + 3
            while not vectors._waiting_queries and time.monotonic() < deadline:
                time.sleep(.005)
            assert vectors._waiting_queries == 1
            pool.submit(vectors.table).result(timeout=2)
        finally:
            release.set()
        indexing.result(timeout=5)
        query.result(timeout=5)
    assert calls == ['document-one', 'question', 'document-two']


def test_legacy_ready_sources_are_requeued_once_without_losing_citations(indexed):
    book, source = indexed
    message = db.add_message(book, 'assistant', 'Supported [1]', [{'text': 'Original quotation'}])
    db.execute('UPDATE sources SET index_version=2 WHERE id=?', (source,))
    db.execute("INSERT OR REPLACE INTO settings VALUES ('rerank','true')")
    db.save_settings({'semantic': True})
    db.initialize()
    row = db.one('SELECT * FROM sources WHERE id=?', (source,))
    assert row['status'] == 'queued' and row['chunks'] == 0
    assert db.one('SELECT citations FROM messages WHERE id=?', (message,))['citations']
    assert not db.one("SELECT * FROM settings WHERE key='rerank'")
    db.execute("UPDATE sources SET status='ready',index_version=? WHERE id=?", (INDEX_VERSION, source))
    db.initialize()
    assert db.one('SELECT status FROM sources WHERE id=?', (source,))['status'] == 'ready'


def test_images_index_without_ocr_and_keep_pixel_provenance(library, monkeypatch, multimodal):
    stream = BytesIO()
    Image.new('RGB', (40, 20), 'red').save(stream, format='PNG')
    source = jobs.add_file(library, 'diagram.png', stream.getvalue())
    db.save_settings({'semantic': True})
    class Splitter:
        def split_text(self, text):
            return [text]
    monkeypatch.setattr(jobs, 'make_splitter', lambda semantic: Splitter())
    recorded = []
    monkeypatch.setattr(vectors, 'insert', lambda rows: recorded.extend(rows))
    jobs.ingest(db.one('SELECT * FROM sources WHERE id=?', (source,)))
    assert len(recorded) == 1 and recorded[0]['image']
    row = db.one('SELECT * FROM chunks WHERE source_id=?', (source,))
    assert row['text'] == '' and row['locator'] == 'Image'
    attachments = answer_images([row])
    with Image.open(BytesIO(attachments[0]['data'])) as image:
        assert image.size == (40, 20) and image.getpixel((0, 0))[0] > 240


def test_rebuild_never_retargets_a_legacy_citation(library):
    source = jobs.add_file(library, 'notes.txt', b'New chunk boundaries and evidence.')
    legacy_id = source + '_0'
    jobs.flush([{'id': legacy_id, 'source_id': source, 'notebook_id': library,
                 'ordinal': 0, 'locator': 'Lines 1', 'text': 'Original citation snapshot.'}], False)
    db.execute('UPDATE sources SET index_version=2,chunks=1 WHERE id=?', (source,))
    jobs.ingest(db.one('SELECT * FROM sources WHERE id=?', (source,)))
    assert not db.one('SELECT id FROM chunks WHERE id=?', (legacy_id,))
    assert db.one('SELECT id FROM chunks WHERE source_id=?', (source,))['id'].startswith(source + f'_g{INDEX_VERSION}_')


def test_image_candidates_remain_distinct_and_scoped(library, monkeypatch, multimodal):
    source = jobs.add_file(library, 'visuals.txt', b'Visual test fixture')
    jobs.flush([{'id': f'{source}_{i}', 'source_id': source, 'notebook_id': library,
                 'ordinal': i, 'locator': f'Page {i}', 'text': '',
                 'image': json.dumps({'path': '/fixture.pdf', 'page': i, 'digest': 'fixture'})}
                for i in range(2)], False)
    db.execute("UPDATE sources SET status='ready' WHERE id=?", (source,))
    db.save_settings({'semantic': True})
    monkeypatch.setattr(vectors, 'search', lambda *args: [
        {'id': f'{source}_{i}', 'semantic_score': .8} for i in range(2)])
    result = retrieve(library, 'Which diagram shows a processor?')
    assert len(result.passages) == 2
    assert all(p['source_id'] == source for p in result.passages)
    assert 'attached image' in chat.evidence_text(result)
    db.execute('UPDATE sources SET selected=0 WHERE id=?', (source,))
    assert not retrieve(library, 'processor').passages


def test_cosine_vector_search_filters_before_limit(library, monkeypatch):
    # Deterministic vectors exercise real LanceDB storage, not model quality.
    def unit(position):
        value = np.zeros(EMBED_DIM, dtype=np.float32)
        value[position] = 1
        return value.tolist()
    vectors.table().add([
        {'id': 'allowed', 'source_id': 'selected', 'notebook_id': library, 'vector': unit(1)},
        {'id': 'private', 'source_id': 'hidden', 'notebook_id': library, 'vector': unit(0)},
        {'id': 'otherbook', 'source_id': 'selected', 'notebook_id': 'other', 'vector': unit(0)}])
    monkeypatch.setattr(vectors, 'query_vector', lambda query: unit(0))
    result = vectors.search('fixture', library, ['selected'], limit=1)
    assert [row['id'] for row in result] == ['allowed']
    assert abs(result[0]['semantic_score']) < 1e-5


def test_text_and_image_gates_use_their_own_similarity_scale(library, monkeypatch, multimodal):
    source = jobs.add_file(library, 'scores.txt', b'Fixture')
    jobs.flush([{'id': 'text', 'source_id': source, 'notebook_id': library, 'ordinal': 0,
                 'locator': 'Text', 'text': 'Archive shelves.'},
                {'id': 'image', 'source_id': source, 'notebook_id': library, 'ordinal': 1,
                 'locator': 'Image', 'text': '', 'image': 'fixture'}], False)
    db.execute("UPDATE sources SET status='ready' WHERE id=?", (source,))
    db.save_settings({'semantic': True})
    monkeypatch.setattr(vectors, 'search', lambda *args: [
        {'id': 'text', 'semantic_score': .63}, {'id': 'image', 'semantic_score': .63}])
    assert [p['id'] for p in retrieve(library, 'processor diagram').passages] == ['image']


def test_cache_separates_text_from_images_and_reuses_identical_pixels(tmp_path):
    class Model:
        calls = []
        def embed(self, inputs, **kwargs):
            self.calls.append(inputs)
            return [np.array([1, i], dtype=np.float32) for i, _ in enumerate(inputs)]
    model = Model()
    image = {'path': 'first.png', 'digest': 'same-pixels', 'frame': 0}
    embed_cached(tmp_path, ['same-pixels', image], model)
    embed_cached(tmp_path, [{**image, 'path': 'duplicate.png'}], model)
    assert len(model.calls) == 1 and len(model.calls[0]) == 2
    embed_cached(tmp_path, [{**image, 'digest': 'changed-pixels'}], model)
    assert len(model.calls) == 2


def test_cache_invalidates_when_model_revision_changes(tmp_path, monkeypatch):
    from local_notebook.retrieval import embedding_cache
    class Model:
        calls = 0
        def embed(self, inputs, **kwargs):
            self.calls += 1
            return [np.array([1], dtype=np.float32) for _ in inputs]
    model = Model()
    embed_cached(tmp_path, ['original'], model)
    monkeypatch.setattr(embedding_cache, 'EMBED_REVISION', 'another-checkpoint')
    embed_cached(tmp_path, ['original'], model)
    assert model.calls == 2


def test_invalid_gemma_vectors_fail_before_indexing():
    import pytest
    class Encoder:
        def encode(self, inputs, **kwargs):
            return np.full((len(inputs), EMBED_DIM), np.nan)
    model = object.__new__(vectors.GemmaEmbedding)
    model.model = Encoder()
    with pytest.raises(ValueError, match='invalid embeddings'):
        model.encode(['fixture'], prompt_name='Document')


def test_visual_pdf_keeps_page_number_without_ocr(tmp_path):
    import pypdfium2 as pdfium
    path = tmp_path / 'scan.pdf'
    Image.new('RGB', (40, 20), 'blue').save(path, 'PDF')
    blocks = list(parse(path, visual=True))
    visuals = [block for block in blocks if block.image]
    assert len(visuals) == 1 and visuals[0].locator == 'Page 1'
    assert visuals[0].image['page'] == 0
    with pdfium.PdfDocument(path) as document:
        assert len(document) == 1


def test_embedded_office_images_decode_from_their_original_source(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    picture = BytesIO()
    Image.new('RGB', (50, 30), 'green').save(picture, 'PNG')
    document = Document()
    document.add_picture(BytesIO(picture.getvalue()))
    document.save(tmp_path / 'notes.docx')
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(
        BytesIO(picture.getvalue()), Inches(1), Inches(1))
    deck.save(tmp_path / 'slides.pptx')
    for name in ('notes.docx', 'slides.pptx'):
        visuals = [block for block in parse(tmp_path / name, visual=True) if block.image]
        assert len(visuals) == 1
        with open_image(visuals[0].image) as decoded:
            assert decoded.size == (50, 30)


def test_context_caps_images_without_breaking_citation_numbers():
    passages = [{'id': str(i), 'name': 'Scan', 'locator': str(i), 'text': '', 'image': 'fixture'}
                for i in range(6)]
    evidence = chat.fit_evidence(Evidence(passages, 'Hybrid retrieval', 0))
    assert len(evidence.passages) == 4 and evidence.warning
    assert chat.cited_passages('Diagram [4]', evidence)[0]['id'] == '3'


def test_compatible_provider_receives_numbered_image_evidence(library, monkeypatch, multimodal):
    db.save_settings({'provider': 'Local / compatible', 'model': 'vision-fixture'})
    received = []
    def handler(request):
        received.append(json.loads(request.content))
        return httpx.Response(200, text='data: {"choices":[{"delta":{"content":"Diagram [2]"}}]}\n\ndata: [DONE]\n\n')
    async def run():
        await providers.close_clients()
        client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        monkeypatch.setattr(providers, 'compatible_client', lambda: client)
        try:
            return await providers.complete('Use evidence', 'Question', images=[
                {'number': 2, 'data': b'fixture-jpeg', 'mime_type': 'image/jpeg'}])
        finally:
            await client.aclose()
    assert asyncio.run(run()) == 'Diagram [2]'
    content = received[0]['messages'][1]['content']
    assert content[1]['text'] == 'Source evidence [2]:'
    assert content[2]['image_url']['url'].startswith('data:image/jpeg;base64,')
