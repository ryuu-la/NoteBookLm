import csv
import json
from contextlib import closing
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from zipfile import ZipFile, is_zipfile

from bs4 import BeautifulSoup


@dataclass
class Block:
    text: str
    locator: str
    progress: float = 0.0
    image: dict | None = None


SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".txt", ".md",
             ".html", ".htm", ".json", ".png", ".jpg", ".jpeg", ".tiff", ".doc", ".ppt", ".xls"}
IMAGE_TYPES = {'.png', '.jpg', '.jpeg', '.tiff'}


def validate_archive(path: Path) -> None:
    if is_zipfile(path):
        with ZipFile(path) as archive:
            items = archive.infolist()
            if len(items) > 20000 or sum(item.file_size for item in items) > 1_000_000_000:
                raise ValueError("The expanded document is too large to process safely.")


def parse(path: Path, *, visual=False, ocr=True):
    extension = path.suffix.lower()
    validate_archive(path)
    if extension == ".pdf":
        yield from pdf_blocks(path, visual=visual, ocr=ocr)
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        total = max(1, len(document.paragraphs) + len(document.tables))
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                yield Block(paragraph.text, f"Paragraph {number}", number / total)
        for number, table in enumerate(document.tables, 1):
            yield Block("\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows),
                        f"Table {number}", (len(document.paragraphs) + number) / total)
        if visual:
            with ZipFile(path) as archive:
                members = [name for name in archive.namelist() if name.startswith('word/media/')
                           and Path(name).suffix.lower() in {'.png', '.jpg', '.jpeg', '.tiff'}]
            for number, member in enumerate(members, 1):
                yield Block('', f"Embedded image {number}", 1, {"path": str(path), "member": member})
    elif extension == ".pptx":
        from pptx import Presentation
        slides = Presentation(path).slides
        for number, slide in enumerate(slides, 1):
            texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
            for shape in slide.shapes:
                if shape.has_table:
                    texts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide:
                texts.append(slide.notes_slide.notes_text_frame.text)
            yield Block("\n".join(texts), f"Slide {number}", number / max(1, len(slides)))
            if visual:
                for shape_number, shape in enumerate(slide.shapes):
                    if shape.shape_type == 13:  # A picture, including its source pixels.
                        yield Block('', f"Slide {number} · image {shape_number + 1}", number / max(1, len(slides)),
                                    {"path": str(path), "slide": number - 1, "shape": shape_number})
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book:
                yield from row_blocks(sheet.iter_rows(values_only=True), sheet.title)
        finally:
            book.close()
    elif extension in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
            yield from row_blocks(csv.reader(handle, delimiter="\t" if extension == ".tsv" else ","), "Rows")
    elif extension in {".png", ".jpg", ".jpeg", ".tiff"}:
        if visual:
            from PIL import Image
            with Image.open(path) as image:
                frames = getattr(image, 'n_frames', 1)
            for frame in range(frames):
                yield Block('', f"Image {frame + 1}" if frames > 1 else "Image", (frame + 1) / frames,
                            {"path": str(path), "frame": frame})
            # OCR is optional; visual indexing works without a second model.
        elif ocr:
            yield Block(ocr_image(str(path)), "Image")
    elif extension in {".doc", ".ppt", ".xls"}:
        from .legacy import convert
        yield from parse(convert(path), visual=visual, ocr=ocr)
    elif extension in SUPPORTED:
        if extension in {".html", ".htm", ".json"}:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            if extension in {".html", ".htm"}:
                soup = BeautifulSoup(text, "html.parser")
                for item in soup(["script", "style", "nav", "footer"]):
                    item.decompose()
                text = soup.get_text("\n", strip=True)
            else:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            for start in range(0, len(text), 16000):
                yield Block(text[start:start + 16000], f"Section {start // 16000 + 1}", min(1, (start + 16000) / max(1, len(text))))
        else:
            with path.open(encoding="utf-8-sig", errors="replace") as handle:
                lines, first, length = [], 1, 0
                for number, line in enumerate(handle, 1):
                    lines.append(line)
                    length += len(line)
                    if length >= 12000:
                        yield Block("".join(lines), f"Lines {first}–{number}")
                        lines, first, length = [], number + 1, 0
                if lines:
                    yield Block("".join(lines), f"Lines {first}–{first + len(lines) - 1}")
    else:
        raise ValueError(f"Unsupported file type: {extension}")


def pdf_blocks(path: Path, *, visual=False, ocr=True):
    """Native PDF text extraction, retaining the Python reader as a fallback."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        yield from pdf_blocks_python(path, visual=visual, ocr=ocr)
        return
    try:
        document = pdfium.PdfDocument(path)
    except pdfium.PdfiumError:
        # The fallback also reports encrypted documents with an actionable error.
        yield from pdf_blocks_python(path, visual=visual, ocr=ocr)
        return
    with document:
        total = len(document)
        for number in range(total):
            with closing(document[number]) as page:
                with closing(page.get_textpage()) as textpage:
                    text = textpage.get_text_bounded().replace("\r\n", "\n")
                has_image = (visual or ocr) and any(page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE]))
                # Page borders and heading underlines should not turn every text
                # page into an expensive visual embedding. Retain complex figures.
                has_visual = visual and (has_image or sum(1 for _ in page.get_objects(
                    filter=[pdfium.raw.FPDF_PAGEOBJ_PATH])) >= 8)
                if ocr and not visual and len(text.strip()) < 20 and has_image:
                    bitmap = page.render(scale=1.5)
                    try:
                        text = ocr_image(bitmap.to_numpy())
                    finally:
                        bitmap.close()
            yield Block(text, f"Page {number + 1}", (number + 1) / total)
            if visual and has_visual:
                yield Block('', f"Page {number + 1}", (number + 1) / total,
                            {"path": str(path), "page": number})


def pdf_blocks_python(path: Path, *, visual=False, ocr=True):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("This PDF is password protected. Import an unlocked copy.")
    total = len(reader.pages)
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if ocr and not visual and len(text.strip()) < 20 and page.images:
            text = ocr_pdf_page(path, number - 1)
        yield Block(text, f"Page {number}", number / total)
        if visual and page.images:
            yield Block('', f"Page {number}", number / total, {"path": str(path), "page": number - 1})


def row_blocks(rows, label: str):
    header, batch, start = "", [], 2
    for number, row in enumerate(rows, 1):
        line = " | ".join(str(value) if value is not None else "" for value in row)
        if number == 1:
            header = line
            continue
        batch.append(line)
        if len(batch) >= 30:
            yield Block(header + "\n" + "\n".join(batch), f"{label} · rows {start}–{number}")
            batch, start = [], number + 1
    if batch:
        yield Block(header + "\n" + "\n".join(batch), f"{label} · rows {start}–{start + len(batch) - 1}")
    elif header and start == 2:
        yield Block(header, f"{label} · row 1")


@lru_cache(maxsize=1)
def ocr_engine():
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise ValueError("Scanned pages need OCR. Install the optional package: pip install -e '.[ocr]'.") from exc
    return RapidOCR()


def ocr_image(image) -> str:
    result, _ = ocr_engine()(image)
    return "\n".join(item[1] for item in (result or []))


def ocr_pdf_page(path: Path, number: int) -> str:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise ValueError("This PDF contains scanned/empty pages. Install the OCR extra to read them.") from exc
    with pdfium.PdfDocument(path) as document:
        return ocr_image(document[number].render(scale=1.5).to_numpy())
