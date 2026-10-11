"""Decode the same source image for embedding and grounded answers."""
import json
from pathlib import Path
from zipfile import ZipFile
from contextlib import closing, contextmanager
from io import BytesIO

from PIL import Image, ImageOps
from ..config import TEXT_ONLY


@contextmanager
def open_image(reference: dict):
    if "page" in reference:
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(reference["path"]) as document:
            with closing(document[reference["page"]]) as page:
                with closing(page.render(scale=1.5)) as bitmap:
                    picture = bitmap.to_pil().convert("RGB")
    elif 'member' in reference:
        with ZipFile(reference['path']) as archive:
            with Image.open(BytesIO(archive.read(reference['member']))) as source:
                picture = ImageOps.exif_transpose(source).convert('RGB')
    elif 'slide' in reference:
        from pptx import Presentation
        slide = Presentation(reference['path']).slides[reference['slide']]
        with Image.open(BytesIO(slide.shapes[reference['shape']].image.blob)) as source:
            picture = ImageOps.exif_transpose(source).convert('RGB')
    else:
        with Image.open(reference["path"]) as source:
            source.seek(reference.get("frame", 0))
            picture = ImageOps.exif_transpose(source).convert("RGB")
    try:
        picture.thumbnail((2048, 2048))
        yield picture
    finally:
        picture.close()


def image_cache_identity(reference: dict) -> dict:
    """Content identity survives duplicate imports under different original paths."""
    if 'digest' not in reference:
        import hashlib
        reference = {**reference, 'digest': hashlib.sha256(Path(reference['path']).read_bytes()).hexdigest()}
    return {key: value for key, value in reference.items() if key != 'path'}


def answer_images(passages: list[dict], limit=4) -> list[dict]:
    """Attach numbered images only after text evidence has been budgeted."""
    if TEXT_ONLY:
        return []
    result = []
    for number, passage in enumerate(passages, 1):
        if not passage.get("image"):
            continue
        with open_image(json.loads(passage["image"])) as picture:
            stream = BytesIO()
            picture.save(stream, format="JPEG", quality=90)
        result.append({"number": number, "data": stream.getvalue(), "mime_type": "image/jpeg"})
        if len(result) == limit:
            break
    return result
